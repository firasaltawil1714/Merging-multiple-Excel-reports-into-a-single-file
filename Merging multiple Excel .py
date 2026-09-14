from pathlib import Path
from openpyxl import Workbook, load_workbook

script_dir = Path(__file__).parent
reports_folder = script_dir / "reports"  
output_path = script_dir / "merged_report.xlsx"


def get_excel_files(folder):
    
    return list(folder.glob("*.xlsx"))


def read_rows(file_path):
    
    workbook = load_workbook(file_path)
    sheet = workbook.active
    return [tuple(cell.value for cell in row) for row in sheet.iter_rows()]

def merge_all(files):

    all_rows = []
    for i, file in enumerate(files):
        rows = read_rows(file)
        if i == 0:
            all_rows.extend(rows)  # Keep header from the first file
        else:
            all_rows.extend(rows[1:])  # Skip header for subsequent files
    return all_rows

def save_merged(rows, path):
    
    workbook = Workbook()
    sheet = workbook.active

    for row in rows:
        sheet.append(row)

    workbook.save(path)

def main():
    if not reports_folder.exists():
        print("Reports folder does not exist.")
        return

    files = get_excel_files(reports_folder)
    if not files:
        print("No Excel files found.")
        return

    all_rows = merge_all(files)
    save_merged(all_rows, output_path)
    print(f"Merged {len(files)} files into {output_path.name}")


if __name__ == "__main__":
    main()